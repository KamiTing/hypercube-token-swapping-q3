import csv
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "output" / "qk_special_cases_routes.xlsx"
csv.field_size_limit(min(sys.maxsize, 2_147_483_647))

SOURCES = {
    "Q8": [
        ROOT / "output" / "q8_case1_trim_20260605_154219" / "qk_special_cases.csv",
        ROOT / "output" / "q8_case2_trim_20260605_154927" / "qk_special_cases.csv",
        ROOT
        / "output"
        / "q4_q11_cuda_all_bw512_pool65536_win65536_restart2048_plateau2048_perturb010_cub_20260613_113055"
        / "qk_special_cases.csv",
    ],
    "Q9": [
        ROOT / "output" / "q9_case1_disk_bw256_20260607_113040" / "qk_special_cases.csv",
        ROOT / "output" / "q9_case2_disk_bw256_path_20260609_170556" / "qk_special_cases.csv",
        ROOT / "output" / "q9_cuda_diverse_bw256_pool4096_win4096_perturb010_20260613_090936" / "qk_special_cases.csv",
        ROOT
        / "output"
        / "q4_q11_cuda_all_bw512_pool65536_win65536_restart2048_plateau2048_perturb010_cub_20260613_113055"
        / "qk_special_cases.csv",
    ],
    "Q10": [
        ROOT / "output" / "cuda_opt_verify_q10_bw512_case1" / "qk_special_cases.csv",
        ROOT / "output" / "cuda_fix_verify_q10_case2" / "qk_special_cases.csv",
        ROOT
        / "output"
        / "q10_cuda_bw512_pool32768_win32768_restart1024_plateau1536_perturb010_case2_20260612_222245"
        / "qk_special_cases.csv",
        ROOT
        / "output"
        / "q4_q11_cuda_all_bw512_pool65536_win65536_restart2048_plateau2048_perturb010_cub_20260613_113055"
        / "qk_special_cases.csv",
    ],
    "Q11": [
        ROOT
        / "output"
        / "q11_cuda_cub_bw512_pool65536_win65536_restart2048_plateau2048_perturb010_path_20260612_231014"
        / "qk_special_cases.csv",
        ROOT
        / "output"
        / "q4_q11_cuda_all_bw512_pool65536_win65536_restart2048_plateau2048_perturb010_cub_20260613_113055"
        / "qk_special_cases.csv",
    ],
    "Q12": [
        ROOT
        / "output"
        / "q12_cuda_gpu_retained_bw1024_pool262144_win65536_restart4096_plateau1536_perturb015_cub_path_20260613_132118"
        / "qk_special_cases.csv",
    ],
}

BLACK = "111827"
SLATE = "334155"
MUTED = "475569"
RED = "EF4444"
BLUE = "2563EB"
TITLE_FILL = "DBEAFE"
HEADER_FILL = "2563EB"
ROUTE_FILL = "F8FAFC"
SUMMARY_FILL = "F1F5F9"
BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def read_rows(path):
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def parse_state(text):
    return [int(x) for x in text.strip().split()]


def parse_path(text):
    if not text or not text.strip():
        return []
    out = []
    for part in text.strip().split():
        u, v = part.split("-", 1)
        out.append((int(u), int(v)))
    return out


def tuple_text(values):
    return "(" + ", ".join(str(v) for v in values) + ")"


def case_number(name):
    marker = "case"
    idx = name.lower().rfind(marker)
    if idx < 0:
        return 999999
    try:
        return int(name[idx + len(marker) :])
    except ValueError:
        return 999999


def route_chunks(swaps, max_chars=28000):
    chunks = []
    current = "Route "
    for u, v in swaps:
        piece = f"({u}, {v})"
        next_text = f"{current}{piece}" if current == "Route " else f"{current}, {piece}"
        if len(next_text) > max_chars and current != "Route ":
            chunks.append(current)
            current = piece
        else:
            current = next_text
    if current and current != "Route ":
        chunks.append(current)
    return chunks or ["Route"]


def solved_rows_for_sheet(sheet_name):
    target_dim = int(sheet_name[1:])
    best_by_case = {}
    for source in SOURCES[sheet_name]:
        if not source.exists():
            continue
        for row in read_rows(source):
            if int(row.get("dim", -1)) != target_dim:
                continue
            if row.get("beam_status") != "solved" or row.get("beam_path_valid") != "1":
                continue
            case = row["case_name"]
            old = best_by_case.get(case)
            key = (int(row["beam_steps"]), float(row["beam_sec"]))
            if old is None or key < (int(old["beam_steps"]), float(old["beam_sec"])):
                best_by_case[case] = row
    rows = list(best_by_case.values())
    rows.sort(key=lambda r: case_number(r["case_name"]))
    return rows


def set_cell(cell, value, font=None, fill=None, align=None, border=True):
    cell.value = value
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if border:
        cell.border = BORDER


def clear_cell(cell):
    cell.value = None
    cell.font = Font(name="Calibri", size=6, color=BLACK)
    cell.fill = PatternFill(fill_type=None)
    cell.alignment = Alignment()
    cell.border = Border()


def unmerge_cell_if_needed(ws, row_idx, col_idx):
    coordinate = ws.cell(row_idx, col_idx).coordinate
    for merged_range in list(ws.merged_cells.ranges):
        if coordinate in merged_range:
            ws.unmerge_cells(str(merged_range))
            return


def clear_sheet_layout(ws):
    for merge in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merge))
    ws.delete_rows(1, ws.max_row)
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 8.43


def write_route_chunks(ws, start_row, start_col, chunks, route_font, route_fill, top_align):
    route_col = start_col + 1
    r = start_row
    for chunk in chunks:
        set_cell(
            ws.cell(r, route_col),
            chunk,
            font=route_font,
            fill=route_fill,
            align=top_align,
        )
        ws.row_dimensions[r].height = 64
        r += 1
    return r


def write_case(ws, row, start_col, dim):
    state = parse_state(row["state_perm"])
    swaps = parse_path(row["beam_path"])
    current = list(state)
    target = list(range(len(current)))

    font_size = 6
    title_font = Font(name="Calibri", size=font_size, bold=True, color="0F172A")
    header_font = Font(name="Calibri", size=font_size, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=font_size, color=BLACK)
    route_font = Font(name="Calibri", size=font_size, italic=True, color=SLATE)
    summary_font = Font(name="Calibri", size=font_size, bold=True, color="0F172A")
    muted_font = Font(name="Calibri", size=font_size, bold=True, color=MUTED)
    title_fill = PatternFill("solid", fgColor=TITLE_FILL)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    route_fill = PatternFill("solid", fgColor=ROUTE_FILL)
    summary_fill = PatternFill("solid", fgColor=SUMMARY_FILL)
    top_align = Alignment(vertical="top", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    end_col = start_col + 2
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    set_cell(
        ws.cell(1, start_col),
        f"Case {case_number(row['case_name'])}: {row['case_name']}",
        font=title_font,
        fill=title_fill,
        align=center_align,
    )
    for col in range(start_col + 1, end_col + 1):
        ws.cell(1, col).border = BORDER
        ws.cell(1, col).fill = title_fill

    headers = ["Packet Route based on the fixed host labeled", "Permutation", "Outcome"]
    for offset, value in enumerate(headers):
        set_cell(
            ws.cell(2, start_col + offset),
            value,
            font=header_font,
            fill=header_fill,
            align=center_align,
        )

    r = 3
    for i, (u, v) in enumerate(swaps):
        before = list(current)
        current[u], current[v] = current[v], current[u]
        after = list(current)
        after_is_target = after == target
        before_color = BLACK if i == 0 else (RED if i % 2 == 1 else BLUE)
        after_color = BLACK if after_is_target else (RED if i % 2 == 0 else BLUE)

        set_cell(ws.cell(r, start_col), f"Swap ({u}, {v})", font=body_font, align=top_align)
        set_cell(
            ws.cell(r, start_col + 1),
            tuple_text(before),
            font=Font(name="Calibri", size=font_size, color=before_color),
            align=top_align,
        )
        set_cell(
            ws.cell(r, start_col + 2),
            tuple_text(after),
            font=Font(name="Calibri", size=font_size, color=after_color),
            align=top_align,
        )
        r += 1

    if current != target:
        raise RuntimeError(f"{row['case_name']} Beam path does not replay to identity")

    chunks = route_chunks(swaps)
    r = write_route_chunks(ws, r, start_col, chunks, route_font, route_fill, top_align)

    summary = [
        [f"Steps: {len(swaps)}", f"Beam: {row['beam_status']} ({row['beam_steps']})", f"Batcher: {'success' if row['batcher_success'] == '1' else 'fail'} ({row['batcher_swaps']})"],
        [f"Strong LB: {row['strong_lb']}", f"Beam path valid: {row['beam_path_valid']}", f"Batcher path valid: {row['batcher_path_valid']}"],
    ]
    for line_idx, line in enumerate(summary):
        for offset, value in enumerate(line):
            set_cell(
                ws.cell(r, start_col + offset),
                value,
                font=summary_font if line_idx == 0 else muted_font,
                fill=summary_fill,
                align=top_align,
            )
        r += 1

    return r - 1


def format_sheet(ws, dim, max_row, case_count):
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 21
    ws.row_dimensions[2].height = 25.5

    widths = {
        8: (28, 174, 174),
        9: (28, 185, 185),
        10: (28, 174, 174),
        11: (28, 205, 205),
        12: (28, 245, 245),
    }[dim]
    for case_idx in range(case_count):
        start_col = 1 + case_idx * 4
        for offset, width in enumerate(widths):
            ws.column_dimensions[get_column_letter(start_col + offset)].width = width
        if case_idx < case_count - 1:
            ws.column_dimensions[get_column_letter(start_col + 3)].width = 3

    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max(3, case_count * 4 - 1)):
        for cell in row:
            if cell.value is None and cell.border == Border():
                continue
            if cell.font:
                copied = copy(cell.font)
                copied.name = "Calibri"
                copied.size = 6
                cell.font = copied


def rebuild_sheet(wb, title, rows, dim, index):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, index)
    clear_sheet_layout(ws)

    max_row = 1
    for case_idx, row in enumerate(rows):
        max_row = max(max_row, write_case(ws, row, 1 + case_idx * 4, dim))

    format_sheet(ws, dim, max_row, len(rows))
    return ws


def center_existing_route_rows(wb):
    for ws in wb.worksheets:
        for start_col in range(1, ws.max_column + 1, 4):
            route_col = start_col + 1
            right_col = start_col + 2
            if right_col > ws.max_column:
                continue
            for row_idx in range(1, ws.max_row + 1):
                source = ws.cell(row_idx, start_col)
                middle = ws.cell(row_idx, route_col)
                right = ws.cell(row_idx, right_col)
                if not (isinstance(source.value, str) and source.value.startswith("Route")):
                    continue
                if middle.value not in (None, "") or right.value not in (None, ""):
                    continue

                unmerge_cell_if_needed(ws, row_idx, start_col)
                unmerge_cell_if_needed(ws, row_idx, route_col)
                unmerge_cell_if_needed(ws, row_idx, right_col)
                source = ws.cell(row_idx, start_col)
                target = ws.cell(row_idx, route_col)
                target.value = source.value
                target.font = copy(source.font)
                target.fill = copy(source.fill)
                target.alignment = copy(source.alignment)
                target.border = copy(source.border)
                clear_cell(source)


def main():
    requested = [arg.upper() for arg in sys.argv[1:]]
    requested = requested or ["Q8", "Q9", "Q10", "Q11"]
    unknown = [sheet for sheet in requested if sheet not in SOURCES]
    if unknown:
        raise RuntimeError(f"Unknown sheet(s): {unknown}")

    wb = load_workbook(WORKBOOK_PATH)

    summaries = {}
    for sheet_name in requested:
        rows = solved_rows_for_sheet(sheet_name)
        expected = 1 if sheet_name == "Q12" else 2
        if len(rows) != expected:
            raise RuntimeError(f"Expected {expected} {sheet_name} rows, found {len(rows)}")

        dim = int(sheet_name[1:])
        previous = f"Q{dim - 1}"
        if sheet_name in wb.sheetnames:
            index = wb.sheetnames.index(sheet_name)
        elif previous in wb.sheetnames:
            index = wb.sheetnames.index(previous) + 1
        else:
            index = len(wb.sheetnames)
        rebuild_sheet(wb, sheet_name, rows, dim, index)
        summaries[sheet_name] = [(r["case_name"], r["beam_steps"]) for r in rows]

    center_existing_route_rows(wb)

    wb.save(WORKBOOK_PATH)
    print(
        {
            "workbook": str(WORKBOOK_PATH),
            **summaries,
        }
    )


if __name__ == "__main__":
    main()
